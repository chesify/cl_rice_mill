import csv
import io

from odoo import api, fields, models
from odoo.exceptions import AccessError, UserError

# column key -> accepted headers (lower case); the template uses the first one
IMPORT_COLUMNS = {
    'farmer_code': ['farmer code', 'namba ya mkulima'],
    'farmer_name': ['farmer name', 'jina la mkulima'],
    'phone': ['phone', 'simu'],
    'village': ['village', 'kijiji'],
    'loan_amount': ['loan', 'mkopo'],
    'tractor_amount': ['tractor', 'trekta'],
    'harvester_amount': ['harvester', 'kivunio'],
    'operator_amount': ['operators paid', 'waendeshaji'],
    'storage_amount': ['storage', 'ghala'],
    'milling_amount': ['milling', 'kukoboa'],
    'other_amount': ['other', 'mengineyo'],
    'paddy_bags': ['paddy bags', 'magunia ya mpunga'],
    'paddy_kg': ['paddy kg', 'kilo za mpunga'],
    'rice_bags': ['rice bags', 'magunia ya mchele'],
    'rice_kg': ['rice kg', 'kilo za mchele'],
}

# debt column -> (company receivable account field, label)
DEBT_TYPES = [
    ('loan_amount', 'cl_rice_loan_account_id', "Loan"),
    ('tractor_amount', 'cl_rice_service_account_id', "Tractor services"),
    ('harvester_amount', 'cl_rice_service_account_id', "Harvester services"),
    ('operator_amount', 'cl_rice_service_account_id', "Operators paid"),
    ('storage_amount', 'cl_rice_service_account_id', "Storage"),
    ('milling_amount', 'cl_rice_service_account_id', "Milling"),
    ('other_amount', 'cl_rice_service_account_id', "Other"),
]


class ClRiceOpeningBalance(models.Model):
    _name = 'cl.rice.opening.balance'
    _description = "Opening Balance Import"
    _inherit = ['mail.thread']
    _order = 'id desc'
    _check_company_auto = True

    name = fields.Char(
        string="Reference", required=True, readonly=True, copy=False, default=lambda self: self.env._("New"))
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company, index=True)
    currency_id = fields.Many2one(related='company_id.currency_id')
    season_id = fields.Many2one(
        'cl.rice.season', string="Season", required=True, check_company=True,
        domain="[('state', '!=', 'closed'), ('company_id', '=', company_id)]")
    date = fields.Date(required=True, default=fields.Date.context_today, help="Accounting date of the opening entries.")
    file = fields.Binary(string="File (XLSX or CSV)", attachment=False)
    file_name = fields.Char()
    state = fields.Selection(
        [('draft', "Draft"), ('validated', "Checked"), ('posted', "Posted"), ('reversed', "Reversed")],
        default='draft', required=True, tracking=True, copy=False)
    line_ids = fields.One2many('cl.rice.opening.balance.line', 'batch_id', string="Farmers")
    line_count = fields.Integer(compute='_compute_totals')
    error_count = fields.Integer(compute='_compute_totals')
    debt_total = fields.Monetary(compute='_compute_totals')
    paddy_kg_total = fields.Float(string="Paddy (kg)", compute='_compute_totals')
    rice_kg_total = fields.Float(string="Rice (kg)", compute='_compute_totals')
    reversal_reason = fields.Text(readonly=True, copy=False)

    @api.depends('line_ids.error', 'line_ids.debt_total', 'line_ids.paddy_kg', 'line_ids.rice_kg')
    def _compute_totals(self):
        for batch in self:
            batch.line_count = len(batch.line_ids)
            batch.error_count = len(batch.line_ids.filtered('error'))
            batch.debt_total = sum(batch.line_ids.mapped('debt_total'))
            batch.paddy_kg_total = sum(batch.line_ids.mapped('paddy_kg'))
            batch.rice_kg_total = sum(batch.line_ids.mapped('rice_kg'))

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            company = self.env['res.company'].browse(vals.get('company_id')) or self.env.company
            if vals.get('name', self.env._("New")) == self.env._("New"):
                vals['name'] = self.env['ir.sequence'].with_company(company).next_by_code('cl.rice.opening.balance') or '/'
        return super().create(vals_list)

    @api.ondelete(at_uninstall=False)
    def _unlink_only_draft(self):
        if self.filtered(lambda b: b.state == 'posted'):
            raise UserError(self.env._("A posted opening balance must be reversed, not deleted."))

    def _cl_check_rights(self):
        if not self.env.user.has_group('cl_rice_mill.cl_rice_mill_group_manager'):
            raise AccessError(self.env._("Only a Rice Mill Manager can import opening balances."))

    # ------------------------------------------------------------------
    # Import
    # ------------------------------------------------------------------
    def action_import_file(self):
        self.ensure_one()
        self._cl_check_rights()
        if self.state not in ('draft', 'validated'):
            raise UserError(self.env._("Only a draft batch can be imported into."))
        if not self.file:
            raise UserError(self.env._("Upload the filled-in template first."))
        rows = self._cl_read_rows(self.file.content, self.file_name or '')
        self.line_ids.unlink()
        self.write({'line_ids': [(0, 0, vals) for vals in rows], 'state': 'draft'})
        return self.action_validate()

    @api.model
    def _cl_read_rows(self, content, file_name):
        if file_name.lower().endswith('.csv'):
            text = content.decode('utf-8-sig')
            table = list(csv.reader(io.StringIO(text)))
        else:
            import openpyxl  # noqa: PLC0415 (only needed for XLSX imports)
            try:
                workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            except Exception as error:  # noqa: BLE001 openpyxl raises many kinds of errors on a bad file
                raise UserError(self.env._("The file could not be read as XLSX or CSV: %s", error)) from error
            sheet = workbook.worksheets[0]
            table = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not table:
            raise UserError(self.env._("The file is empty."))
        header_map = {}
        for index, header in enumerate(table[0]):
            label = str(header or '').strip().lower()
            for key, aliases in IMPORT_COLUMNS.items():
                if label in aliases:
                    header_map[index] = key
        if 'farmer_name' not in header_map.values() and 'farmer_code' not in header_map.values():
            raise UserError(self.env._("The file has no farmer code or farmer name column: use the template."))
        rows = []
        for sequence, row in enumerate(table[1:], start=1):
            if not any(cell not in (None, '') for cell in row):
                continue
            vals = {'sequence': sequence}
            for index, key in header_map.items():
                cell = row[index] if index < len(row) else None
                if key in ('farmer_code', 'farmer_name', 'phone', 'village'):
                    vals[key] = str(cell).strip() if cell not in (None, '') else False
                else:
                    try:
                        vals[key] = float(str(cell).replace(',', '')) if cell not in (None, '') else 0.0
                    except ValueError:
                        vals[key] = 0.0
                        vals['error'] = self.env._("Column %(column)s is not a number: %(value)s", column=key, value=cell)
            rows.append(vals)
        return rows

    def action_download_template(self):
        return {
            'type': 'ir.actions.act_url',
            'url': '/rice/opening_balance/template',
            'target': 'self',
        }

    @api.model
    def _cl_template_xlsx(self):
        import xlsxwriter  # noqa: PLC0415 (only needed to build the template)
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        sheet = workbook.add_worksheet("Opening Balances")
        bold = workbook.add_format({'bold': True, 'bg_color': '#E8F5E9', 'border': 1})
        headers = [aliases[0].title() for aliases in IMPORT_COLUMNS.values()]
        for column, header in enumerate(headers):
            sheet.write(0, column, header, bold)
            sheet.set_column(column, column, max(len(header) + 2, 12))
        example = ['F00001', "Juma Mwakalinga", '255712345678', "Igurusi",
                   500000, 200000, 0, 40000, 0, 0, 0, 20, 1800, 0, 0]
        for column, value in enumerate(example):
            sheet.write(1, column, value)
        workbook.close()
        return output.getvalue()

    # ------------------------------------------------------------------
    # Validation (preview)
    # ------------------------------------------------------------------
    def action_validate(self):
        self.ensure_one()
        self._cl_check_rights()
        if self.state not in ('draft', 'validated'):
            raise UserError(self.env._("Only a draft batch can be checked."))
        if not self.line_ids:
            raise UserError(self.env._("There is nothing to check: import a file or add farmers."))
        self.line_ids._cl_validate()
        self.state = 'draft' if self.error_count else 'validated'
        return True

    # ------------------------------------------------------------------
    # Posting
    # ------------------------------------------------------------------
    def action_post(self):
        self.ensure_one()
        self._cl_check_rights()
        self.action_validate()
        if self.error_count:
            raise UserError(self.env._("Fix the %s lines with errors before posting.", self.error_count))
        company = self.company_id
        if not company.cl_rice_loan_account_id or not company.cl_rice_settlement_journal_id:
            raise UserError(self.env._("Run the Rice Mill setup first: the farmer accounts are missing."))
        for line in self.line_ids:
            line._cl_post()
        self.state = 'posted'
        self.message_post(body=self.env._(
            "Opening balances of %(count)s farmers posted by %(user)s.", count=len(self.line_ids), user=self.env.user.name))

    def action_open_reverse(self):
        self.ensure_one()
        self._cl_check_rights()
        return {
            'name': self.env._("Reverse Opening Balances"),
            'type': 'ir.actions.act_window',
            'res_model': 'cl.rice.reason',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_res_model': self._name, 'default_res_id': self.id, 'default_action': 'reverse'},
        }

    def _cl_reason_actions(self):
        return {'reverse': self._cl_reverse}

    def _cl_reverse(self, reason, amount=None):
        """Reverse the whole batch: opening entries and opening stock."""
        self.ensure_one()
        self._cl_check_rights()
        if self.state != 'posted':
            raise UserError(self.env._("Only a posted batch can be reversed."))
        locked = self.line_ids.farmer_season_id.filtered('is_locked')
        if locked:
            raise UserError(self.env._("Farmer season %s is already settled: reverse its settlement first.", locked[0].name))
        for line in self.line_ids:
            line._cl_reverse(reason)
        self.write({'state': 'reversed', 'reversal_reason': reason})
        self.message_post(body=self.env._(
            "Opening balances reversed by %(user)s. Reason: %(reason)s", user=self.env.user.name, reason=reason))
